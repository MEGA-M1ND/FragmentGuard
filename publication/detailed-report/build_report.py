"""Render the detailed research document. No experiment execution or model calls.

Requires pandoc, python-docx, ReportLab and DejaVu fonts in the build environment.
Run: python3 publication/detailed-report/build_report.py
"""

import importlib.util
import json
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
STEM = "FragmentGuard-detailed-report"
SOURCE = HERE / (STEM + ".md")
INK = "#142b3b"
TEAL = "#176b80"


def word():
    output = HERE / (STEM + ".docx")
    subprocess.run(
        [
            "pandoc",
            SOURCE.name,
            "-f",
            "gfm-tex_math_dollars",
            "-o",
            str(output),
            "--resource-path",
            str(HERE),
        ],
        cwd=HERE,
        check=True,
    )
    doc = Document(output)
    styles = {s.name: s for s in doc.styles}
    section = doc.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    section.different_first_page_header_footer = True
    for name in ["Normal", "Body Text", "First Paragraph"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string("142B3B")
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(7)
    for name, size in [("Heading 1", 30), ("Heading 2", 16), ("Heading 3", 12)]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("176B80" if name == "Heading 3" else "142B3B")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    if "Source Code" in doc.styles:
        styles["Source Code"].font.name = "Consolas"
        styles["Source Code"].font.size = Pt(9)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name == "Source Code":
            p.paragraph_format.keep_together = True
        if p.text in ["Contents", "1. Executive summary"]:
            p.paragraph_format.page_break_before = True
        if p.text == "Resource provenance and oversight of multi-agent AI workflows":
            p.style = styles["Subtitle"]
        if p.text.startswith("Figure "):
            p.style = styles["Caption"]
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = Inches(6.55)
        shape.height = int(shape.width * ratio)
    for table in doc.tables:
        table.style = styles.get("Table Grid", styles.get("Table", table.style))
        for n, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trpr.append(cant)
            if n == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "142B3B" if n == 0 else ("F1F6F8" if n % 2 else "FFFFFF"))
                cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = n < len(table.rows) - 1
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if n == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
    header = section.header.paragraphs[0]
    header.text = "FRAGMENTGUARD  |  DETAILED RESEARCH REPORT"
    header.style = styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.text = "29 October 2026  •  github.com/MEGA-M1ND/FragmentGuard  |  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.title = "FragmentGuard: resource provenance and oversight of multi-agent AI workflows"
    doc.core_properties.author = "FragmentGuard project"
    doc.core_properties.subject = "Detailed report of the completed research; 29 October 2026"
    doc.save(output)


def pdf():
    # Reuse the tested inline Markdown renderer from the concise publication.
    spec = importlib.util.spec_from_file_location("publication_pdf", ROOT / "scripts/publication_pdf.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    inline = module.inlines
    fonts = Path("/usr/share/fonts/truetype/dejavu")
    for name, file in [
        ("FG", "DejaVuSans.ttf"),
        ("FGBold", "DejaVuSans-Bold.ttf"),
        ("FGItalic", "DejaVuSans-Oblique.ttf"),
        ("FGBoldItalic", "DejaVuSans-BoldOblique.ttf"),
        ("FGMono", "DejaVuSansMono.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / file)))
    pdfmetrics.registerFontFamily("FG", normal="FG", bold="FGBold", italic="FGItalic", boldItalic="FGBoldItalic")
    body = ParagraphStyle(
        "Body",
        fontName="FG",
        fontSize=10,
        leading=14.5,
        textColor=colors.HexColor(INK),
        spaceAfter=7,
        allowWidows=0,
        allowOrphans=0,
    )
    title = ParagraphStyle("Title", parent=body, fontName="FGBold", fontSize=34, leading=40, spaceAfter=15)
    subtitle = ParagraphStyle("Subtitle", parent=body, fontSize=20, leading=27, spaceAfter=20)
    heading = ParagraphStyle(
        "Heading",
        parent=body,
        fontName="FGBold",
        fontSize=15,
        leading=20,
        spaceBefore=16,
        spaceAfter=9,
        keepWithNext=True,
    )
    subhead = ParagraphStyle(
        "Subhead", parent=heading, fontSize=11.5, leading=16, textColor=colors.HexColor(TEAL), spaceBefore=11
    )
    small = ParagraphStyle("Small", parent=body, fontSize=8.5, leading=12, textColor=colors.HexColor("#526473"))
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.7, leading=12, spaceAfter=0)
    cell_head = ParagraphStyle("CellHead", parent=cell, fontName="FGBold", textColor=colors.white)
    code = ParagraphStyle(
        "Code",
        parent=body,
        fontName="FGMono",
        fontSize=8,
        leading=11,
        backColor=colors.HexColor("#f1f5f7"),
        borderPadding=7,
        spaceBefore=5,
        spaceAfter=10,
    )
    width = A4[0] - 108

    class ReportDoc(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name == "Heading":
                title_text = flowable.getPlainText()
                match = re.match(r"^(\d+)\.", title_text)
                if match:
                    key = "section" + match.group(1)
                    self.canv.bookmarkPage(key)
                    self.canv.addOutlineEntry(title_text, key, 0)

    doc = ReportDoc(
        str(HERE / (STEM + ".pdf")),
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=48,
        bottomMargin=48,
        title="FragmentGuard: detailed research report",
        author="FragmentGuard project",
    )
    blocks = json.loads(
        subprocess.check_output(
            [
                "pandoc",
                str(SOURCE),
                "-f",
                "gfm-tex_math_dollars",
                "-t",
                "json",
            ]
        )
    )["blocks"]
    story = [Spacer(1, 40), Paragraph("COMPLETED RESEARCH · DETAILED REPORT", small), Spacer(1, 15)]
    i = 0
    while i < len(blocks):
        block = blocks[i]
        tag, value = block["t"], block.get("c")
        if tag == "Header":
            label = inline(value[2])
            if label in ["Contents", "1. Executive summary"]:
                story.append(PageBreak())
            style = title if value[0] == 1 else (heading if value[0] == 2 else subhead)
            if label == "Resource provenance and oversight of multi-agent AI workflows":
                style = subtitle
            story.append(Paragraph(label, style))
        elif tag in ["Para", "Plain"]:
            if len(value) == 1 and value[0]["t"] == "Image":
                path = (HERE / value[0]["c"][2][0]).resolve()
                img = Image(str(path))
                parts = [Image(str(path), width=width, height=width * img.imageHeight / img.imageWidth)]
                if i + 1 < len(blocks) and blocks[i + 1]["t"] == "Para":
                    parts.append(Paragraph(inline(blocks[i + 1]["c"]), small))
                    i += 1
                story.extend([KeepTogether(parts), Spacer(1, 8)])
            else:
                story.append(Paragraph(inline(value), body))
        elif tag in ["OrderedList", "BulletList"]:
            entries = value[1] if tag == "OrderedList" else value
            start = value[0][0] if tag == "OrderedList" else 1
            for n, entry in enumerate(entries, start=start):
                prefix = f"{n}. " if tag == "OrderedList" else "• "
                for sub in entry:
                    label = prefix + inline(sub["c"])
                    if tag == "OrderedList" and len(entries) == 16:
                        label = f'<a color="{TEAL}" href="#section{n}">{label}</a>'
                    story.append(Paragraph(label, body))
        elif tag == "CodeBlock":
            parts = [Preformatted(value[1], code, maxLineLength=85, splitChars=" /")]
            if story and isinstance(story[-1], Paragraph):
                parts.insert(0, story.pop())
            while story and isinstance(story[-1], Paragraph) and getattr(story[-1].style, "keepWithNext", False):
                parts.insert(0, story.pop())
            story.append(KeepTogether(parts))
        elif tag == "Table":
            headers = value[3][1]
            rows = list(headers)
            for tb in value[4]:
                rows.extend(tb[2])
                rows.extend(tb[3])
            data = [
                [
                    Paragraph("<br/>".join(inline(b["c"]) for b in c[4]), cell_head if n < len(headers) else cell)
                    for c in row[1]
                ]
                for n, row in enumerate(rows)
            ]
            cols = len(data[0])
            fractions = [0.36, 0.64] if cols == 2 else [0.36, 0.32, 0.32]
            table = Table(
                data,
                colWidths=[width * f for f in fractions],
                repeatRows=len(headers),
                hAlign="LEFT",
                spaceBefore=6,
                spaceAfter=12,
            )
            style = [
                ("BACKGROUND", (0, 0), (-1, len(headers) - 1), colors.HexColor(INK)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#d5e0e5")),
            ]
            for n in range(len(headers), len(data)):
                if n % 2:
                    style.append(("BACKGROUND", (0, n), (-1, n), colors.HexColor("#f1f6f8")))
            table.setStyle(TableStyle(style))
            parts = [table]
            if story and isinstance(story[-1], Paragraph) and len(story[-1].getPlainText()) < 350:
                parts.insert(0, story.pop())
            while story and isinstance(story[-1], Paragraph) and getattr(story[-1].style, "keepWithNext", False):
                parts.insert(0, story.pop())
            story.append(KeepTogether(parts))
        else:
            raise ValueError("Unsupported Markdown block: " + tag)
        i += 1

    def page(canvas, current):
        if current.page == 1:
            return
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#526473"))
        canvas.setFont("FG", 7.5)
        canvas.drawString(54, A4[1] - 27, "FRAGMENTGUARD · DETAILED RESEARCH REPORT")
        canvas.drawString(54, 25, "29 October 2026 · github.com/MEGA-M1ND/FragmentGuard")
        canvas.drawRightString(A4[0] - 54, 25, str(current.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=page, onLaterPages=page)


if __name__ == "__main__":
    word()
    pdf()
    print("Built DOCX and PDF in", HERE)
